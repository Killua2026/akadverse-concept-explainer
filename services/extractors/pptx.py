"""PPTX slide text and existing speaker notes."""

from pptx import Presentation

from .errors import ExtractionError


def extract_text(file_path: str) -> str:
    try:
        parts = []
        for slide in Presentation(file_path).slides:
            parts.extend(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None and frame.text.strip():
                    parts.append(frame.text)
        content = "\n".join(parts)
    except Exception as error:
        raise ExtractionError("Unable to extract text from PPTX") from error
    if not content.strip():
        raise ExtractionError("No extractable text in PPTX")
    return content
