"""DOCX paragraphs followed by table cell text."""

from docx import Document

from .errors import ExtractionError


def extract_text(file_path: str) -> str:
    try:
        document = Document(file_path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        content = "\n".join(parts)
    except Exception as error:
        raise ExtractionError("Unable to extract text from DOCX") from error
    if not content.strip():
        raise ExtractionError("No extractable text in DOCX")
    return content
