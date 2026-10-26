import asyncio
import sys

if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())

from dotenv import load_dotenv
load_dotenv()

from contextlib import asynccontextmanager
from pathlib import Path
from uuid import uuid4

import pytest


@pytest.fixture
def storage_session():
    @asynccontextmanager
    async def isolated_session():
        from sqlalchemy import text
        from sqlalchemy.ext.asyncio import async_sessionmaker

        from config import StorageSettings
        from db import Database

        database = Database(StorageSettings.from_env())
        schema = f"concept_explainer_test_{uuid4().hex}"
        migrations_dir = Path(__file__).resolve().parents[1] / "migrations"
        migrations = []
        for name in ("001_pgvector_storage.sql", "002_jobs_table.sql"):
            migrations.append(await asyncio.to_thread(
                (migrations_dir / name).read_text, encoding="utf-8"
            ))
        migration = "\n".join(migrations)
        statements = [
            statement.strip()
            for statement in migration.replace("concept_explainer", schema).split(";")
            if statement.strip() and statement.strip() not in {"BEGIN", "COMMIT"}
        ]
        try:
            async with database.engine.connect() as connection:
                transaction = await connection.begin()
                try:
                    extension = await connection.scalar(text(
                        "SELECT extname FROM pg_extension WHERE extname = 'vector'"
                    ))
                    assert extension == "vector", "Apply the approved migration first."
                    # Exercise reruns using only a generated test schema.
                    for _ in range(2):
                        for statement in statements:
                            await connection.execute(text(statement))
                    sessions = async_sessionmaker(
                        connection,
                        expire_on_commit=False,
                        join_transaction_mode="create_savepoint",
                    )
                    async with sessions() as session:
                        yield session, schema
                finally:
                    await transaction.rollback()
        finally:
            await database.dispose()

    return isolated_session


@pytest.fixture
def worker_database(storage_session, monkeypatch):
    @asynccontextmanager
    async def isolated_worker():
        from types import SimpleNamespace
        from sqlalchemy import text
        from sqlalchemy.ext.asyncio import async_sessionmaker
        from workers import ingestion

        async with storage_session() as (session, schema):
            monkeypatch.setattr(ingestion, "text", lambda sql: text(sql.replace("concept_explainer", schema)))
            yield SimpleNamespace(sessions=async_sessionmaker(
                session.bind, expire_on_commit=False,
                join_transaction_mode="create_savepoint",
            )), schema
    return isolated_worker


@pytest.fixture
def office_document(tmp_path):
    """Generate small Office files without adding binary fixtures."""
    def create(kind, paragraph="", table="", notes=""):
        path = tmp_path / f"lecture.{kind}"
        if kind == "docx":
            from docx import Document
            document = Document()
            if paragraph:
                document.add_paragraph(paragraph)
            if table:
                document.add_table(rows=1, cols=1).cell(0, 0).text = table
            document.save(path)
        else:
            from pptx import Presentation
            from pptx.util import Inches
            document = Presentation()
            if paragraph or notes:
                slide = document.slides.add_slide(document.slide_layouts[6])
                if paragraph:
                    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = paragraph
                if notes:
                    slide.notes_slide.notes_text_frame.text = notes
            document.save(path)
        return path
    return create
